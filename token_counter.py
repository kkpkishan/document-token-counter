#!/usr/bin/env python3
import os
import csv
import sys
import argparse
import logging
import platform
from pathlib import Path
from typing import List, Dict, Tuple, Optional
from concurrent.futures import ProcessPoolExecutor, as_completed

from tqdm import tqdm
import tiktoken

logger = logging.getLogger("token_counter")

# -----------------------------
# tokenization
# -----------------------------
def get_encoder(encoding: Optional[str], model: Optional[str]):
    try:
        if model:
            return tiktoken.encoding_for_model(model)
    except Exception as e:
        logger.warning(f"Could not load encoding for model '{model}': {e}. Falling back...")
    try:
        if encoding:
            return tiktoken.get_encoding(encoding)
    except Exception as e:
        logger.warning(f"Could not load encoding '{encoding}': {e}. Falling back...")
    return tiktoken.get_encoding("cl100k_base")

def count_tokens(text: str, encoder) -> int:
    try:
        return len(encoder.encode(text or ""))
    except Exception as e:
        logger.error(f"Error counting tokens: {e}")
        return 0

# -----------------------------
# extractors
# -----------------------------
def process_pdf(file_path: str) -> str:
    try:
        from pypdf import PdfReader
        with open(file_path, "rb") as fh:
            reader = PdfReader(fh)
            parts = []
            for page in reader.pages:
                try:
                    parts.append(page.extract_text() or "")
                except Exception:
                    parts.append("")
            return "\n".join(parts)
    except Exception as e:
        logger.error(f"Error processing PDF file {file_path}: {e}")
        return ""

def process_docx(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text]
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        logger.error(f"Error processing DOCX file {file_path}: {e}")
        return ""

def process_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error processing TXT file {file_path}: {e}")
            return ""
    except Exception as e:
        logger.error(f"Error processing TXT file {file_path}: {e}")
        return ""

def process_pptx(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        parts: List[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
            if hasattr(slide, "has_notes_slide") and slide.has_notes_slide and slide.notes_slide:
                notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                if notes_frame and notes_frame.text:
                    parts.append(notes_frame.text)
        return "\n".join(parts)
    except Exception as e:
        logger.error(f"Error processing PPTX file {file_path}: {e}")
        return ""

def process_excel(file_path: str, excel_max_rows: Optional[int]) -> str:
    try:
        import pandas as pd
        suffix = Path(file_path).suffix.lower()
        if suffix == ".xls":
            xls = pd.read_excel(file_path, sheet_name=None, engine="xlrd", dtype=str, nrows=excel_max_rows)
        else:
            xls = pd.read_excel(file_path, sheet_name=None, engine="openpyxl", dtype=str, nrows=excel_max_rows)
        frames = list(xls.values())
        if not frames:
            return ""
        df = pd.concat(frames, ignore_index=True)
        return df.to_csv(index=False, header=False) or ""
    except Exception as e:
        logger.error(f"Error processing Excel file {file_path}: {e}")
        return ""

def process_csv(file_path: str, csv_max_rows: Optional[int]) -> str:
    try:
        import pandas as pd
        df = pd.read_csv(file_path, dtype=str, nrows=csv_max_rows)
        return df.to_csv(index=False, header=False) or ""
    except Exception as e:
        logger.error(f"Error processing CSV file {file_path}: {e}")
        return ""

def process_doc_ppt_via_com(file_path: str, app: str) -> str:
    if platform.system() != "Windows":
        logger.warning(f"Skipping {file_path}: requires Windows/MS Office.")
        return ""
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except Exception:
        logger.warning(f"Skipping {file_path}: pywin32 not installed or COM unavailable.")
        return ""
    try:
        pythoncom.CoInitialize()
        if app == "Word":
            word = win32com.client.Dispatch("Word.Application")
            doc = word.Documents.Open(str(file_path))
            text = doc.Content.Text
            doc.Close(False)
            word.Quit()
            return text
        elif app == "PowerPoint":
            ppt = win32com.client.Dispatch("PowerPoint.Application")
            pres = ppt.Presentations.Open(str(file_path), ReadOnly=True)
            parts = []
            for slide in pres.Slides:
                for shape in slide.Shapes:
                    if getattr(shape, "HasTextFrame", False):
                        try:
                            parts.append(shape.TextFrame.TextRange.Text)
                        except Exception:
                            pass
            pres.Close()
            ppt.Quit()
            return "\n".join(parts)
        return ""
    except Exception as e:
        logger.error(f"Error processing {Path(file_path).suffix} file {file_path} via COM: {e}")
        return ""
    finally:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

# -----------------------------
# worker & orchestration
# -----------------------------
# globals for workers (initialized in init_worker)
_ENCODER = None
_EXCEL_MAX = None
_CSV_MAX = None

def init_worker(encoding: Optional[str], model: Optional[str], excel_max_rows: Optional[int], csv_max_rows: Optional[int]):
    global _ENCODER, _EXCEL_MAX, _CSV_MAX
    _ENCODER = get_encoder(encoding, model)
    _EXCEL_MAX = excel_max_rows
    _CSV_MAX = csv_max_rows

def process_one_file(file_path: str) -> Dict[str, str]:
    """Runs in worker process. Returns {'file_path': str, 'token_count': str}."""
    try:
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            text = process_pdf(file_path)
        elif ext == ".docx":
            text = process_docx(file_path)
        elif ext == ".txt":
            text = process_txt(file_path)
        elif ext == ".pptx":
            text = process_pptx(file_path)
        elif ext in (".xls", ".xlsx"):
            text = process_excel(file_path, _EXCEL_MAX)
        elif ext == ".csv":
            text = process_csv(file_path, _CSV_MAX)
        else:
            # unsupported here (or handled serially)
            return {"file_path": file_path, "token_count": "0"}
        tokens = count_tokens(text, _ENCODER)
        return {"file_path": file_path, "token_count": str(tokens)}
    except Exception as e:
        logger.error(f"Unhandled error in worker for {file_path}: {e}")
        return {"file_path": file_path, "token_count": "0"}

def iter_files(root: str) -> List[str]:
    files = []
    for dirpath, _, filenames in os.walk(root):
        for name in filenames:
            files.append(os.path.join(dirpath, name))
    return files

def create_csv_report(results: List[Dict[str, str]], total_tokens: int, total_files: int, output_path: str):
    try:
        with open(output_path, "w", newline="", encoding="utf-8") as csvfile:
            fieldnames = ["file_path", "token_count"]
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()
            for row in results:
                writer.writerow(row)
            writer.writerow({"file_path": "Total", "token_count": str(total_tokens)})
            writer.writerow({"file_path": "Total Files", "token_count": str(total_files)})
        logger.info(f"Report saved to: {output_path}")
    except Exception as e:
        logger.error(f"Error creating CSV report: {e}")

def main():
    parser = argparse.ArgumentParser(description="Count tokens in files across multiple folders.")
    parser.add_argument("folders", nargs="+", help="Full path(s) to source folders to process")
    parser.add_argument("-o", "--output", required=True, help="Full path to output CSV file")

    enc = parser.add_argument_group("Tokenization")
    enc.add_argument("--model", default=None, help="Model name to select encoding (e.g., gpt-4o, gpt-4o-mini)")
    enc.add_argument("--encoding", default=None, help="Explicit tiktoken encoding (e.g., cl100k_base, o200k_base)")

    limits = parser.add_argument_group("Row limits")
    limits.add_argument("--excel-max-rows", type=int, default=None, help="Max rows per sheet to read from Excel")
    limits.add_argument("--csv-max-rows", type=int, default=None, help="Max rows to read from CSV")

    perf = parser.add_argument_group("Performance")
    perf.add_argument("--workers", type=int, default=os.cpu_count() or 1,
                      help="Number of parallel worker processes (default: CPU count)")
    perf.add_argument("--quiet", action="store_true", help="Reduce logging verbosity")

    args = parser.parse_args()

    logging.basicConfig(
        level=logging.WARNING if args.quiet else logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )

    output_path = args.output if args.output.lower().endswith(".csv") else f"{args.output}.csv"

    logger.info("Token Counter")
    logger.info("Input folders:")
    for folder in args.folders:
        logger.info(f"- {folder}")
    logger.info(f"Output file: {output_path}")
    logger.info(f"Workers: {args.workers}")

    # list files and split into parallelizable vs COM-legacy
    all_files: List[str] = []
    for folder in args.folders:
        all_files.extend(iter_files(folder))

    # legacy COM types kept serial to avoid Windows COM issues
    serial_exts = {".doc", ".ppt"}
    serial_files = [p for p in all_files if Path(p).suffix.lower() in serial_exts]
    parallel_files = [p for p in all_files if p not in serial_files]

    # results aggregation
    results: List[Dict[str, str]] = []
    total_tokens = 0
    total_files = 0

    # 1) process COM files serially (if any)
    if serial_files:
        logger.info(f"Processing {len(serial_files)} COM file(s) serially")
        for fp in tqdm(serial_files, desc="COM files", unit="file"):
            ext = Path(fp).suffix.lower()
            text = process_doc_ppt_via_com(fp, app="Word" if ext == ".doc" else "PowerPoint")
            tokens = count_tokens(text, get_encoder(args.encoding, args.model))
            results.append({"file_path": fp, "token_count": str(tokens)})
            total_tokens += tokens
            total_files += 1

    # 2) process everything else in parallel
    if parallel_files:
        logger.info(f"Processing {len(parallel_files)} file(s) in parallel")
        with ProcessPoolExecutor(max_workers=max(1, args.workers),
                                 initializer=init_worker,
                                 initargs=(args.encoding, args.model, args.excel_max_rows, args.csv_max_rows)) as ex:
            futures = {ex.submit(process_one_file, fp): fp for fp in parallel_files}
            for fut in tqdm(as_completed(futures), total=len(futures), desc="Parallel", unit="file"):
                row = fut.result()
                results.append(row)
                total_files += 1
                try:
                    total_tokens += int(row["token_count"])
                except Exception:
                    pass

    create_csv_report(results, total_tokens, total_files, output_path)

    logger.info("Processing complete!")
    logger.info(f"Total tokens: {total_tokens}")
    logger.info(f"Total files processed: {total_files}")

if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        logger.warning("Interrupted by user. Exiting...")
        sys.exit(130)
